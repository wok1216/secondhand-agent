import hashlib
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


DATA_DIR = Path(__file__).resolve().parent.parent / "data"


HEADER_ALIASES = {
    "id": {
        "id",
        "매물id",
        "상품id",
        "제품id",
        "listingid",
        "productid",
    },

    "category": {
        "category",
        "카테고리",
        "대분류",
        "상품카테고리",
        "품목",
    },

    "subcategory": {
        "subcategory",
        "서브카테고리",
        "소분류",
        "세부카테고리",
    },

    "title": {
        "title",
        "상품명",
        "제품명",
        "제목",
        "상품제목",
        "매물명",
    },

    "description": {
        "description",
        "desc",
        "설명",
        "판매글",
        "본문",
        "내용",
        "상세내용",
        "상세설명",
    },

    "price": {
        "price",
        "가격",
        "판매가",
        "판매가격",
        "금액",
    },

    "location": {
        "location",
        "거래위치",
        "거래장소",
        "거래지역",
        "지역",
        "위치",
        "주소",
    },

    "distance_km": {
        "distance",
        "distancekm",
        "거리",
        "거리km",
    },
}


STANDARD_FIELDS = set(
    HEADER_ALIASES.keys()
)


def _normalize_header(value):
    if value is None:
        return ""

    return re.sub(
        r"[\s_\-()/]",
        "",
        str(value).strip().lower(),
    )


def _build_alias_map():
    result = {}

    for key, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            result[
                _normalize_header(alias)
            ] = key

    return result


ALIAS_MAP = _build_alias_map()


def _resolve_header(value):
    normalized = _normalize_header(
        value
    )

    return ALIAS_MAP.get(
        normalized
    )


def _clean_value(value):
    if value is None:
        return None

    if isinstance(
        value,
        (datetime, date)
    ):
        return value.isoformat()

    if isinstance(value, str):
        value = value.strip()

        if not value:
            return None

    return value


def _parse_price(value):
    value = _clean_value(
        value
    )

    if value is None:
        return None

    if isinstance(
        value,
        (int, float)
    ):
        return int(value)

    text = str(value).replace(
        ",",
        ""
    ).strip()

    man_match = re.search(
        r"(\d+(?:\.\d+)?)\s*만",
        text
    )

    if man_match:
        return int(
            float(
                man_match.group(1)
            )
            * 10000
        )

    thousand_match = re.search(
        r"(\d+(?:\.\d+)?)\s*천",
        text
    )

    if thousand_match:
        return int(
            float(
                thousand_match.group(1)
            )
            * 1000
        )

    number_match = re.search(
        r"\d+(?:\.\d+)?",
        text
    )

    if not number_match:
        return None

    return int(
        float(
            number_match.group(0)
        )
    )


def _parse_distance_km(value):
    value = _clean_value(
        value
    )

    if value is None:
        return None

    if isinstance(
        value,
        (int, float)
    ):
        return float(value)

    text = str(value).lower()

    km_match = re.search(
        r"(\d+(?:\.\d+)?)\s*km",
        text
    )

    if km_match:
        return float(
            km_match.group(1)
        )

    meter_match = re.search(
        r"(\d+(?:\.\d+)?)\s*m",
        text
    )

    if meter_match:
        return (
            float(
                meter_match.group(1)
            )
            / 1000
        )

    return None


def _detect_header_row(
    worksheet,
    max_scan_rows=15,
):
    """
    처음 몇 행을 살펴보고
    title + price가 존재하는 행을
    매물 데이터 헤더로 판단한다.
    """

    best = None

    max_row = min(
        worksheet.max_row,
        max_scan_rows
    )

    for row_number in range(
        1,
        max_row + 1
    ):
        mapped = {}

        for column_number in range(
            1,
            worksheet.max_column + 1
        ):
            value = worksheet.cell(
                row=row_number,
                column=column_number,
            ).value

            key = _resolve_header(
                value
            )

            if key and key not in mapped:
                mapped[key] = column_number

        # 최소 조건:
        # 상품명 + 가격
        if (
            "title" not in mapped
            or "price" not in mapped
        ):
            continue

        score = len(mapped)

        if (
            best is None
            or score > best["score"]
        ):
            best = {
                "row": row_number,
                "columns": mapped,
                "score": score,
            }

    return best


def _make_generated_id(
    source_file,
    sheet_name,
    row_number,
    title,
):
    raw = (
        f"{source_file}|"
        f"{sheet_name}|"
        f"{row_number}|"
        f"{title}"
    )

    digest = hashlib.sha1(
        raw.encode("utf-8")
    ).hexdigest()[:12]

    return f"ITEM_{digest}"


def _split_category(
    category,
    subcategory,
):
    category = _clean_value(
        category
    )

    subcategory = _clean_value(
        subcategory
    )

    if (
        category
        and not subcategory
    ):
        parts = re.split(
            r"\s*[/>\|]\s*",
            str(category),
            maxsplit=1,
        )

        if len(parts) == 2:
            category = (
                parts[0].strip()
                or None
            )

            subcategory = (
                parts[1].strip()
                or None
            )

    return (
        category or "기타",
        subcategory,
    )


def _read_listing_sheet(
    path,
    worksheet,
    header_info,
):
    header_row = header_info[
        "row"
    ]

    column_map = header_info[
        "columns"
    ]

    # 원본 헤더명 보관
    raw_headers = {}

    for column_number in range(
        1,
        worksheet.max_column + 1
    ):
        header = _clean_value(
            worksheet.cell(
                row=header_row,
                column=column_number,
            ).value
        )

        if header is not None:
            raw_headers[
                column_number
            ] = str(header)

    items = []

    for row_number in range(
        header_row + 1,
        worksheet.max_row + 1
    ):
        row_values = {}

        for key, column_number in (
            column_map.items()
        ):
            row_values[key] = (
                _clean_value(
                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                    ).value
                )
            )

        title = row_values.get(
            "title"
        )

        if not title:
            continue

        price = _parse_price(
            row_values.get(
                "price"
            )
        )

        if price is None:
            continue

        category, subcategory = (
            _split_category(
                row_values.get(
                    "category"
                ),
                row_values.get(
                    "subcategory"
                ),
            )
        )

        location = row_values.get(
            "location"
        )

        distance_km = (
            _parse_distance_km(
                row_values.get(
                    "distance_km"
                )
            )
        )

        # 거리 전용 컬럼이 없으면
        # location 문자열에서 찾아본다.
        if distance_km is None:
            distance_km = (
                _parse_distance_km(
                    location
                )
            )

        metadata = {}

        mapped_columns = set(
            column_map.values()
        )

        for column_number, header in (
            raw_headers.items()
        ):
            if (
                column_number
                in mapped_columns
            ):
                continue

            value = _clean_value(
                worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).value
            )

            if value is None:
                continue

            metadata[
                header
            ] = value

        item_id = row_values.get(
            "id"
        )

        if item_id is None:
            item_id = (
                _make_generated_id(
                    source_file=path.name,
                    sheet_name=worksheet.title,
                    row_number=row_number,
                    title=title,
                )
            )

        item = {
            "id": str(item_id),
            "category": category,
            "subcategory": subcategory,
            "title": str(title),
            "description": (
                str(
                    row_values.get(
                        "description"
                    )
                    or ""
                )
            ),
            "price": price,
            "location": (
                str(location)
                if location
                else ""
            ),
            "distance_km":
                distance_km,

            # 제품별 제각각인 데이터
            "metadata":
                metadata,

            # 디버깅 / 연동 추적용
            "source": {
                "file": path.name,
                "sheet":
                    worksheet.title,
                "row":
                    row_number,
            },
        }

        items.append(
            item
        )

    return items


def _read_workbook(path):
    workbook = load_workbook(
        path,
        data_only=True,
    )

    items = []

    for worksheet in workbook.worksheets:
        header_info = (
            _detect_header_row(
                worksheet
            )
        )

        # 매물 데이터 형식이 아닌 시트
        # 자동 무시
        if header_info is None:
            continue

        items.extend(
            _read_listing_sheet(
                path=path,
                worksheet=worksheet,
                header_info=header_info,
            )
        )

    return items


def get_catalog_items():
    """
    backend/data 안의 모든 xlsx를 읽고
    건지니 공통 매물 스키마로 변환한다.
    """

    items = []

    excel_files = sorted(
        path
        for path in DATA_DIR.glob(
            "*.xlsx"
        )
        if not path.name.startswith(
            "~$"
        )
    )

    for path in excel_files:
        items.extend(
            _read_workbook(
                path
            )
        )

    return items


def get_catalog_item_by_id(
    item_id: str
):
    for item in get_catalog_items():
        if (
            str(item["id"])
            == str(item_id)
        ):
            return item

    return None