import time
from collections import Counter

import requests
from openpyxl import load_workbook


EXCEL_PATH = "data/list.xlsx"
SHEET_NAME = "평가정답_입력금지"
BASE_URL = "http://127.0.0.1:8000"

# 상위 3개만 비교
TARGET_IDS = [
    "NOTEBOOK_001",
    "NOTEBOOK_002",
    "NOTEBOOK_003",
]

TEST_CONDITIONS = {
    "purpose": "대학교 과제 수행 및 코딩",
    "preferred_os": "Windows",
    "preferred_brands": ["삼성", "LG"],
    "max_weight_kg": 1.3,
    "max_price": 500000,
    "max_distance_km": 3.0,
    "preferred_design": "단순한 디자인",
    "preferred_exterior_condition": "양호",
    "max_usage_months": 36,
    "require_normal_function": True,
    "require_charger": True,
}


def load_ground_truth():
    workbook = load_workbook(
        EXCEL_PATH,
        data_only=True
    )

    sheet = workbook[SHEET_NAME]
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]

    result = []

    for row in rows[1:]:
        if not row or not row[0]:
            continue

        data = dict(zip(headers, row))

        # 상위 3개만 남김
        if data["ID"] in TARGET_IDS:
            result.append(data)

    workbook.close()

    # TARGET_IDS 순서대로 정렬
    result.sort(
        key=lambda x: TARGET_IDS.index(x["ID"])
    )

    return result


def predict(item_id):
    response = requests.post(
        f"{BASE_URL}/api/items/{item_id}/evaluate",
        json={
            "conditions": TEST_CONDITIONS
        },
        timeout=120
    )

    response.raise_for_status()

    return response.json()


def main():
    ground_truth = load_ground_truth()

    print()
    print("=" * 60)
    print("건지니 상위 3개 자동 평가")
    print("=" * 60)
    print()

    total = 0
    correct = 0

    results = []

    expected_counter = Counter()
    predicted_counter = Counter()

    for index, truth in enumerate(
        ground_truth,
        start=1
    ):
        item_id = truth["ID"]
        expected = truth["기대분류"]

        print(
            f"[{index}/{len(ground_truth)}] "
            f"{item_id} 평가 중..."
        )

        try:
            prediction = predict(item_id)

            final = prediction["final"]
            predicted = final["verdict"]
            is_correct = (predicted == expected)

            total += 1
            if is_correct:
                correct += 1

            expected_counter[expected] += 1
            predicted_counter[predicted] += 1

            results.append({
                "id": item_id,
                "expected": expected,
                "predicted": predicted,
                "correct": is_correct,
                "summary": final.get("summary", ""),
                "fallback": prediction.get("fallback", False),
                "cache_hit": prediction.get("cache_hit", False),
            })

            print(f"  기대: {expected}")
            print(f"  결과: {predicted}")
            print(f"  판정: {'일치' if is_correct else '불일치'}")

            if prediction.get("cache_hit"):
                print("  캐시 사용")
            if prediction.get("fallback"):
                print("  fallback 결과")

            # 캐시가 아니면 잠깐 쉬기
            if not prediction.get("cache_hit", False):
                time.sleep(8)

        except Exception as e:
            print(f"  오류: {e}")

        print()

    print("=" * 60)
    print("최종 결과")
    print("=" * 60)

    accuracy = (correct / total * 100) if total else 0

    print(f"평가 성공: {total}개")
    print(f"일치: {correct}개")
    print(f"불일치: {total - correct}개")
    print(f"분류 일치율: {accuracy:.1f}%")
    print()

    print("상세 결과")
    for result in results:
        mark = "O" if result["correct"] else "X"
        print(
            f"[{mark}] {result['id']} | "
            f"기대={result['expected']} | "
            f"건지니={result['predicted']}"
        )

        if not result["correct"]:
            print(f"    요약: {result['summary']}")


if __name__ == "__main__":
    main()